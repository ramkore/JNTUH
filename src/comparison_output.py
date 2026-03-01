"""
Output generators for comparison results - Excel and PDF.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XlImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RlImage
from reportlab.lib.enums import TA_CENTER
import os
import uuid

# Path to the header image
HEADER_IMAGE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "header.png")


# ---------------------------------------------------------------------------
#  Excel comparison output
# ---------------------------------------------------------------------------

def generate_comparison_excel(df_comparison, summary, output_path,
                              subject="PPS", department="CIVIL ENGINEERING",
                              year="2025-2026", year_sem="I-I Sem"):
    """Generate an Excel file with color-coded comparison results."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    # Styles
    header_font = Font(name='Calibri', bold=True, size=11)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    yellow_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    header_fill = PatternFill(start_color='DAEEF3', end_color='DAEEF3', fill_type='solid')

    # Header image
    if os.path.exists(HEADER_IMAGE_PATH):
        try:
            img = XlImage(HEADER_IMAGE_PATH)
            img.width = 700
            img.height = 80
            ws.add_image(img, 'A1')
            start_row = 6
        except Exception:
            start_row = 1
    else:
        start_row = 1

    # Title rows
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=13)
    title_cell = ws.cell(row=start_row, column=1, value="MID-I vs MID-II Comparison")
    title_cell.font = Font(name='Calibri', bold=True, size=14)
    title_cell.alignment = center

    start_row += 1
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=13)
    sub_cell = ws.cell(row=start_row, column=1,
                       value=f"Subject: {subject}  |  Dept: {department}  |  Year: {year}  |  {year_sem}")
    sub_cell.font = Font(name='Calibri', size=10)
    sub_cell.alignment = center

    start_row += 2

    # Column headers
    headers = [
        'S.No.', 'Roll No', 'Student Name',
        'Mid-I SA', 'Mid-I Essay', 'Mid-I Assg', 'Mid-I Total',
        'Mid-II SA', 'Mid-II Essay', 'Mid-II Assg', 'Mid-II Total',
        'Difference', 'Status'
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = header_font
        cell.alignment = center
        cell.border = thin_border
        cell.fill = header_fill

    # Column widths
    widths = [6, 18, 22, 8, 8, 8, 10, 8, 8, 8, 10, 10, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Data rows — use enumerate for robust row numbering
    data_start = start_row + 1
    for row_num, (idx, row) in enumerate(df_comparison.iterrows()):
        r = data_start + row_num
        values = [
            row_num + 1, row['Roll No'], row['Student Name'],
            row['Mid1_SA'], row['Mid1_Essay'], row['Mid1_Assignment'], row['Mid1_Total'],
            row['Mid2_SA'], row['Mid2_Essay'], row['Mid2_Assignment'], row['Mid2_Total'],
            row['Difference'], row['Status']
        ]
        # Determine fill color
        status = str(row['Status'])
        fill = None
        if status == 'Improved':
            fill = green_fill
        elif status == 'Declined':
            fill = red_fill
        elif 'Missing' in status:
            fill = yellow_fill

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col_idx, value=val)
            cell.alignment = center
            cell.border = thin_border
            if fill:
                cell.fill = fill

    # Summary section
    summary_row = data_start + len(df_comparison) + 2
    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
    ws.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True, size=12)

    summary_items = [
        ("Total Students", summary['total']),
        ("Improved", summary['improved']),
        ("Declined", summary['declined']),
        ("Same", summary['same']),
        ("Missing in MID-I", summary['missing_mid1']),
        ("Missing in MID-II", summary['missing_mid2']),
    ]
    for i, (label, val) in enumerate(summary_items):
        r = summary_row + 1 + i
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=val)

    wb.save(output_path)


# ---------------------------------------------------------------------------
#  PDF comparison output
# ---------------------------------------------------------------------------

def generate_comparison_pdf(df_comparison, summary, output_path,
                            subject="PPS", department="CIVIL ENGINEERING",
                            year="2025-2026", year_sem="I-I Sem"):
    """Generate a PDF report with comparison table and summary."""
    doc = SimpleDocTemplate(output_path, pagesize=landscape(A4),
                            leftMargin=0.5*inch, rightMargin=0.5*inch,
                            topMargin=0.4*inch, bottomMargin=0.4*inch)
    elements = []
    styles = getSampleStyleSheet()

    # Use unique style names to avoid conflicts on repeated calls
    uid = uuid.uuid4().hex[:6]

    # Header image
    if os.path.exists(HEADER_IMAGE_PATH):
        try:
            img = RlImage(HEADER_IMAGE_PATH)
            img_w = 7.5 * inch
            img_h = 0.85 * inch
            img.drawWidth = img_w
            img.drawHeight = img_h
            elements.append(img)
            elements.append(Spacer(1, 0.15 * inch))
        except Exception:
            pass

    # Title
    title_style = ParagraphStyle(f'CompTitle_{uid}', parent=styles['Title'], fontSize=14,
                                 alignment=TA_CENTER, spaceAfter=4)
    sub_style = ParagraphStyle(f'CompSub_{uid}', parent=styles['Normal'], fontSize=9,
                               alignment=TA_CENTER, spaceAfter=8)

    elements.append(Paragraph("MID-I vs MID-II Comparison", title_style))
    elements.append(Paragraph(
        f"Subject: {subject} &nbsp;|&nbsp; Dept: {department} &nbsp;|&nbsp; Year: {year} &nbsp;|&nbsp; {year_sem}",
        sub_style
    ))

    # Build table data
    header_labels = ['S.No.', 'Roll No', 'Name',
               'M1 SA', 'M1 Esy', 'M1 Asg', 'M1 Tot',
               'M2 SA', 'M2 Esy', 'M2 Asg', 'M2 Tot',
               'Diff', 'Status']

    cell_style = ParagraphStyle(f'Cell_{uid}', parent=styles['Normal'], fontSize=7,
                                alignment=TA_CENTER, leading=9)
    name_style = ParagraphStyle(f'NameCell_{uid}', parent=styles['Normal'], fontSize=7,
                                leading=9)
    hdr_style = ParagraphStyle(f'H_{uid}', parent=styles['Normal'],
                               fontSize=7, alignment=TA_CENTER, leading=9,
                               textColor=colors.black)

    table_data = [[Paragraph(h, hdr_style) for h in header_labels]]

    for row_num, (idx, row) in enumerate(df_comparison.iterrows()):
        name_text = str(row['Student Name'])[:20]  # truncate long names
        diff_val = row['Difference']
        if isinstance(diff_val, float):
            diff_val = round(diff_val, 1)
        table_data.append([
            Paragraph(str(row_num + 1), cell_style),
            Paragraph(str(row['Roll No']), cell_style),
            Paragraph(name_text, name_style),
            Paragraph(str(row['Mid1_SA']), cell_style),
            Paragraph(str(row['Mid1_Essay']), cell_style),
            Paragraph(str(row['Mid1_Assignment']), cell_style),
            Paragraph(str(row['Mid1_Total']), cell_style),
            Paragraph(str(row['Mid2_SA']), cell_style),
            Paragraph(str(row['Mid2_Essay']), cell_style),
            Paragraph(str(row['Mid2_Assignment']), cell_style),
            Paragraph(str(row['Mid2_Total']), cell_style),
            Paragraph(str(diff_val), cell_style),
            Paragraph(str(row['Status']), cell_style),
        ])

    col_widths = [0.35*inch, 1.0*inch, 1.1*inch,
                  0.4*inch, 0.4*inch, 0.4*inch, 0.5*inch,
                  0.4*inch, 0.4*inch, 0.4*inch, 0.5*inch,
                  0.5*inch, 0.9*inch]

    table = Table(table_data, colWidths=col_widths, repeatRows=1)

    # Base style
    style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DAEEF3')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
    ]

    # Color-code rows by status
    for row_num, (idx, row) in enumerate(df_comparison.iterrows()):
        r = row_num + 1  # +1 for header
        status = str(row['Status'])
        if status == 'Improved':
            style_cmds.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#C6EFCE')))
        elif status == 'Declined':
            style_cmds.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#FFC7CE')))
        elif 'Missing' in status:
            style_cmds.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#FFEB9C')))

    table.setStyle(TableStyle(style_cmds))
    elements.append(table)

    # Summary section
    elements.append(Spacer(1, 0.3 * inch))
    sum_title_style = ParagraphStyle(f'SumTitle_{uid}', parent=styles['Heading2'],
                                     fontSize=12, spaceAfter=6)
    elements.append(Paragraph("Summary", sum_title_style))

    summary_data = [
        ['Total Students', str(summary['total'])],
        ['Improved', str(summary['improved'])],
        ['Declined', str(summary['declined'])],
        ['Same', str(summary['same'])],
        ['Missing in MID-I', str(summary['missing_mid1'])],
        ['Missing in MID-II', str(summary['missing_mid2'])],
    ]

    sum_table = Table(summary_data, colWidths=[2 * inch, 1 * inch])
    sum_table.setStyle(TableStyle([
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#DAEEF3')),
    ]))
    elements.append(sum_table)

    doc.build(elements)
